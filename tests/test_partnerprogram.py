"""Partnerprogrammet: produktbiblioteket, Glasrikets program, materialets regler och partneröversikten."""
import re
from pathlib import Path

import golfholes.partnerprogram as pp
from golfholes.partnerprogram import Program, logg
from golfholes.render import partner_deck as D
from golfholes.render import partner_dok as K
from golfholes.render.halkort import ATTRIBUTION, Halkort

ROOT = Path(__file__).parent.parent
GLAS = ROOT / "klubbar" / "glasriket"
pp.MARKERA = Halkort.from_yaml          # utan ortofoto: omarkerad hålbild räcker för texttesterna


def _prog() -> Program:
    return Program.from_dir(GLAS)


def _text(html: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>|data:[^'\"]+", " ", html))


def test_glasrikets_program_och_priser():
    p = _prog()
    assert p.sasong == 2027 and len(p.hal) == 18
    assert {h.nr for h in p.hal if h.premium} == {1, 17}
    assert all(h.pris == (25000 if h.premium else 20000) for h in p.hal)
    assert {m["nyckel"]: m["summa"] for m in p.modeller()} == {"differentierad": 370000, "enhetlig": 360000,
                                                          "enhetlig_25": 450000}
    assert p.helt_program()["summa"] == 502500
    assert [r["summa"] for r in p.scenarier()][:1] == [70000]          # dagens läge: 7 × 10 000
    assert p.befintliga_trappa() == [15000, 17500, 20000]
    assert 6 <= len(p.aktiva()) <= 10                                   # briefen: 6–10 aktiva produkter
    assert p.produkter["korthal"]["namn"] == "Pay & Play-partner"
    assert p.produkter["klubbhusskarm"]["status"] == "vilande"
    assert p.produkter["digital_kanal"]["status"] == "framtid"


def test_ingen_yta_saljs_tva_ganger_och_veckorna_racker():
    p = _prog()
    v = p.partnerveckor()
    assert not v["over"] and v["plan"][30][0] == "huvudpartner"
    assert len(set(v["plan"].values())) == len(v["plan"])            # en partner per vecka
    assert not [k for k in p.kontroller() if k[0] == "stopp"]
    # utan bekräftad skärm lovas ingen skärm i något hålsponsorskap
    for pid in ("halsponsor", "premiumhal"):
        ytor = p.produkter[pid]["ytor"]
        assert not any("skärm" in y or "entré" in y for y in ytor["fysiska"] + ytor["digitala"])
    assert not any(y.get("kraver") for y in p.veckoytor())             # partnerveckan är digital för Glasriket


def test_foretagsdagen_raknas_ur_mallen():
    fd = _prog().foretagsdag(12)
    assert fd["sjalvkostnad"] == 9900 and fd["pris"] == 13500


def test_materialet_foljer_reglerna():
    p = _prog()
    sp_ex, sp_neutral = D.Sponsor(p, True), D.Sponsor(p, False)
    extern = "".join(D.deck(p, sp_ex, "25 oktober 2026")[0] + D.deck(p, sp_neutral, "25 oktober 2026")[0]
                     + D.prospekt_hal(p, p.hal[13], sp_ex)[0] + [D.framtida_kanal(p, sp_neutral, False)]
                     + [K.prislista(p), K.rapport(p, {"foretag": "Nordvik", "hal": 14}, {"mall": True})])
    t = _text(extern)
    assert not re.search(r"GameBook|Emtec|\bAI\b|exkl\. moms|designstudio", t)
    assert "Nordvik-hålet" in t and "[Ert företag]" in t
    assert ATTRIBUTION in extern
    assert "personer ser" not in t                                      # exponering = visningar, inte personer
    assert "Golf GameBook" in D.framtida_kanal(p, sp_ex, True)         # namnet bara i den interna versionen
    # tusentalsavgränsare med mellanslag
    assert "20 000 kr" in t and "20000 kr" not in t


def test_rapporten_hittar_aldrig_pa_varden():
    p = _prog()
    tom = _text(K.rapport(p, {"foretag": "Nordvik", "hal": 14}, {}))
    assert tom.count("fylls i") >= 10 and "MALL" not in tom
    ifylld = _text(K.rapport(p, {"foretag": "Nordvik", "hal": 14},
                             {"matt": [{"vad": "Hålets sida, sidvisningar", "varde": "1 234", "kalla": "Plausible"}]}))
    assert "1 234" in ifylld and "Uppmätt" in ifylld and "Uppskattning" in ifylld


def test_partneroversikten_skapas_och_lases(tmp_path):
    from openpyxl import load_workbook

    p = _prog()
    f = logg.skapa(p, tmp_path / "oversikt.xlsx")
    wb = load_workbook(f)
    assert wb.sheetnames == ["Instruktion", "Partners", "Kontakter", "Leveranser", "Mätvärden", "Ronder"]
    ws = wb["Partners"]
    assert ws.max_row - 1 == 18 + 5                                     # 18 hål + range, putting, P&P, junior, huvud
    rad = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "hal-14")
    rad[3].value, rad[4].value = "såld", "Nordvik"
    wb["Mätvärden"].append(["hal-14", "Sidvisningar", 1234, "uppmätt", "Plausible", "apr–okt"])
    wb["Leveranser"].append(["hal-14", "Sponsorplattan uppsatt", "2027-04-20", None, "Banpersonalen"])
    wb.save(f)
    d = logg.las(f)
    assert d["status"]["hal-14"] == "såld" and d["partners"]["hal-14"]["foretag"] == "Nordvik"
    assert d["rapport"]["hal-14"]["matt"][0]["varde"] == "1 234"
    assert d["rapport"]["hal-14"]["levererat"][0]["varde"] == "2027-04-20"
    assert d["ronder"] is None                                          # inga ronder ifyllda = ingen uppskattning


def test_matvarde_utan_kalla_stoppas(tmp_path):
    from openpyxl import load_workbook

    f = logg.skapa(_prog(), tmp_path / "o.xlsx")
    wb = load_workbook(f)
    wb["Mätvärden"].append(["hal-01", "Sidvisningar", 10, "uppmätt", None, None])
    wb.save(f)
    try:
        logg.las(f)
    except SystemExit as err:
        assert "källa saknas" in str(err)
    else:
        raise AssertionError("ett mätvärde utan källa ska stoppas")


if __name__ == "__main__":
    import tempfile

    for namn, fn in list(globals().items()):
        if namn.startswith("test_"):
            if "tmp_path" in fn.__code__.co_varnames[:fn.__code__.co_argcount]:
                with tempfile.TemporaryDirectory() as d:
                    fn(Path(d))
            else:
                fn()
            print("ok", namn)
